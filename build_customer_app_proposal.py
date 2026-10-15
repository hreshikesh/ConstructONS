from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

out = Path(r'D:\ConstructONS\ConstructONS_Customer_App_Proposal.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.68); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
NAVY='17324D'; ORANGE='E8753D'; PALE='EEF3F7'; LIGHT='F7F9FB'; GRAY='596777'; WHITE='FFFFFF'; GREEN='2B6B57'
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.3); styles['Normal'].font.color.rgb=RGBColor.from_string('253443')
styles['Normal'].paragraph_format.space_after=Pt(5); styles['Normal'].paragraph_format.line_spacing=1.08
for nm,size,color in [('Title',28,NAVY),('Heading 1',17,NAVY),('Heading 2',12,NAVY),('Heading 3',10,ORANGE)]:
 s=styles[nm]; s.font.name='Aptos Display' if nm!='Normal' else 'Aptos'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(9 if nm!='Title' else 0); s.paragraph_format.space_after=Pt(4)

def shade(cell, fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell, text, bold=False, color=None, size=8.4):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
 r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size)
 if color: r.font.color.rgb=RGBColor.from_string(color)
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for i,h in enumerate(headers): shade(t.rows[0].cells[i],NAVY); set_cell_text(t.rows[0].cells[i],h,True,WHITE,8.3)
 for ridx,row in enumerate(rows):
  cells=t.add_row().cells
  for i,val in enumerate(row):
   shade(cells[i], LIGHT if ridx%2==0 else WHITE); set_cell_text(cells[i],str(val),False,None,8.1)
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def bullet(text, level=0):
 p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p
def numbered(text):
 p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p
def callout(label, text):
 t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,PALE); c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
 r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(ORANGE); r.font.size=Pt(9)
 r=p.add_run(text); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(NAVY)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)

# Header / footer
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hr=header.add_run('CONSTRUCTONS  /  CUSTOMER APP PROPOSAL'); hr.bold=True; hr.font.name='Aptos'; hr.font.size=Pt(7.5); hr.font.color.rgb=RGBColor.from_string(GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fr=footer.add_run('Product proposal  •  13 October 2026'); fr.font.size=Pt(7.5); fr.font.color.rgb=RGBColor.from_string(GRAY)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(28); r=p.add_run('PRODUCT PROPOSAL'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(ORANGE)
p=doc.add_paragraph(style='Title'); p.add_run('ConstructONS Customer App')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14); r=p.add_run('A phased plan for project delivery, home planning, and AI-enabled construction guidance'); r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(GRAY)
callout('PROPOSED DIRECTION','Build a customer-facing construction companion. Start with project management in Phase 1, then add home planning, Vastu guidance, budgeting, design exploration, and AI assistance in later phases.')
doc.add_heading('Executive summary', level=1)
doc.add_paragraph('ConstructONS should give customers a single place to manage the journey from a signed project through construction, while also helping prospective homeowners plan what to build. The app is not just a milestone tracker: project management is the first release, and the broader product grows into practical planning and decision-support tools.')
doc.add_paragraph('Phase 1 focuses on a dependable customer project experience: sign-in, project overview, milestones, updates, documents, and a clear route to contact the team. The existing website remains the staff workspace. Staff continue to use its admin login and admin panel to create projects, update milestones, and publish customer-facing information.')

doc.add_heading('Product goals', level=1)
for x in ['Make project progress, next steps, and shared files easy for customers to find.','Help future customers turn plot, family, style, and budget needs into a useful planning brief.','Use AI to explain and explore options, while leaving professional design and construction decisions to qualified people.','Reuse the existing website and backend where practical, so the app does not create a second source of project data.'] : bullet(x)

doc.add_heading('Product boundary and users', level=1)
table(['Area','Customer app','Existing website'],[
('Authentication','Customer sign-in, account, and access to linked projects.','Admin/staff login stays on the website.'),
('Project operations','Customers view updates, milestones, and authorized documents.','Staff create projects, manage records, update stages, and publish updates.'),
('Planning and AI','Customer tools for planning, exploration, and requests for professional review.','Staff review incoming briefs and quote/consultation requests.')],[1.55,2.75,2.75])
callout('DELIVERY ASSUMPTION','Start with a mobile-first responsive app or installable PWA that connects to the current backend. Decide on native iOS/Android apps only after customer usage and device needs are understood.')

doc.add_heading('Phased roadmap', level=1)
table(['Phase','Theme','Customer outcome'],[
('1','Project management','A signed-in customer can follow their project, see progress and updates, access shared files, and contact the team.'),
('2','AI Home Planner','A prospective customer can describe their plot and needs, compare planning options, and save a structured home brief.'),
('3','AI Vastu Guide','Customers can explore common Vastu considerations as an optional preference while comparing layout concepts.'),
('4','Budget and materials','Customers can explore preliminary cost ranges, package choices, materials, finishes, and saved selections.'),
('5','AI design and construction guide','Customers can create inspiration boards and get answers grounded in approved ConstructONS information.'),
('6','Connected customer journey','Consultation and quote requests, approvals, appointments, and deeper project tools connect the planning journey to delivery.')],[.55,1.55,4.95])
p=doc.add_paragraph('Phases are ordered by product dependency, not calendar duration. Delivery estimates should be set after confirming platform choice, team capacity, integrations, and Phase 1 acceptance criteria.')
p.runs[0].italic=True; p.runs[0].font.color.rgb=RGBColor.from_string(GRAY); p.runs[0].font.size=Pt(8.3)

doc.add_page_break()
doc.add_heading('Phase 1 in detail  Project management', level=1)
doc.add_paragraph('Phase 1 establishes the app as the customer’s project companion. It adapts the project and milestone capabilities already present in the website/backend into a complete, mobile-friendly customer experience. Staff continue managing project data on the website.')
doc.add_heading('Phase 1 objectives', level=2)
for x in ['Give each customer a secure view of only their linked project or projects.','Make current stage, overall progress, recent news, and upcoming work immediately visible.','Keep stage history, notes, and progress photos together in one timeline.','Make project documents and support contact easy to reach.','Use the current website admin workflow as the source of truth for project updates.'] : bullet(x)
doc.add_heading('Customer-facing capabilities', level=2)
table(['Capability','Phase 1 behavior'],[
('Customer access','Sign in using the customer account flow; provide sign-out and session expiry handling. Admin/staff authentication is not added to the app.'),
('Project home','Show project name, location or agreed identifier, current status, overall progress, current milestone, latest update, and next expected step.'),
('Milestone timeline','Show all project stages with status, dates, customer-visible notes, and completion progress where appropriate.'),
('Progress updates','Present dated updates and photos in a readable history; distinguish new updates from older activity.'),
('Documents','List and open only files shared with that customer, such as approved proposals, drawings, invoices, or project documents.'),
('Contact and help','Provide a direct route to submit a question or request help, including enough project context for the team to respond.'),
('Notifications','In-app notification or email when a new customer-visible milestone or update is published; notification preferences can follow.')],[1.35,5.7])
doc.add_heading('Staff workflow on the website', level=2)
for x in ['Create or select the customer project in the existing admin panel.','Update milestone status, dates, progress notes, and customer-visible photos.','Share documents with the right customer/project.','Review customer questions and follow up through the team’s agreed support channel.'] : numbered(x)

doc.add_heading('Phase 1 delivery slices', level=2)
table(['Slice','Work included','Completion checkpoint'],[
('1A  Foundation','Confirm app shell, responsive navigation, customer session behavior, project ownership rules, and API contracts.','Customer can sign in and only receives their authorized project data.'),
('1B  Project overview','Project summary, progress indicator, current and next milestone, and recent update.','Customer can understand project status from the landing screen.'),
('1C  Timeline and updates','Ten-stage tracker, stage detail, dated notes, and progress photos.','Customer can follow stage changes and see the update history.'),
('1D  Documents and contact','Authorized document list/view, support request/contact path, and helpful empty/error states.','Customer can find shared information and ask for help.'),
('1E  Release readiness','Accessibility and mobile review, security/access review, monitoring, and launch checklist.','Phase 1 acceptance criteria are met and release owner approves launch.')],[1.0,2.9,3.15])
doc.add_heading('Phase 1 acceptance criteria', level=2)
for x in ['A customer can sign in, sign out, and return to their project without exposing admin functionality.','A customer can view only projects and files explicitly linked to their account.','The project home clearly shows the current status, progress, latest update, and next step when available.','All configured milestones display accurate states and customer-visible details from the backend.','A newly published update or milestone change appears in the app and triggers the agreed notification behavior.','The app handles no-project, missing-document, loading, expired-session, and service-error states clearly.','The experience works on common phone screen sizes and remains usable with keyboard and assistive technology.'] : bullet(x)

doc.add_heading('Out of scope for Phase 1', level=2)
for x in ['Admin/staff login or project editing inside the customer app.','AI-generated floor plans, Vastu analysis, or construction decisions.','Payments, contract signing, or binding approvals.','Native app-store release before confirming that a PWA/responsive app is insufficient.'] : bullet(x)

doc.add_heading('Later phases at a glance', level=1)
doc.add_heading('Phase 2  AI Home Planner', level=2)
doc.add_paragraph('Collect plot dimensions, location, family needs, room preferences, number of floors, lifestyle needs, and budget range. Use those inputs to create a saved planning brief and compare early concept options. Include clear assumptions and invite professional review before a concept is used for design or construction.')
doc.add_heading('Phase 3  AI Vastu Guide', level=2)
doc.add_paragraph('Let customers optionally provide plot and entrance orientation, then explore common Vastu preferences in plain language. Explain trade-offs and let customers compare alternatives. Present the feature as informational guidance, not a substitute for architecture, engineering, or local code review.')
doc.add_heading('Phase 4  Budget and Materials', level=2)
doc.add_paragraph('Offer preliminary cost ranges with visible assumptions, package comparisons, material and finish discovery, and saved selections. Costs should be labeled as estimates until reviewed and confirmed by ConstructONS.')
doc.add_heading('Phase 5  AI Design and Construction Guide', level=2)
doc.add_paragraph('Support room inspiration and mood boards, plus a question-answer assistant grounded in approved ConstructONS content. The assistant should say when it is uncertain and offer a human handoff. Generated visuals are inspiration, not construction drawings.')
doc.add_heading('Phase 6  Connected Journey', level=2)
doc.add_paragraph('Connect saved planning briefs and selections to consultation requests and quote workflows. Consider customer approvals, appointment scheduling, and expanded project collaboration after the core app and team processes are stable.')

doc.add_heading('Key product and technical principles', level=1)
for x in ['One source of truth: project records remain in the existing backend and staff administer them through the website.','Least-privilege access: server-side authorization checks every project, file, and conversation against the signed-in customer.','AI with boundaries: AI suggestions are labeled, assumptions are visible, and consequential decisions require professional review.','Progressive delivery: validate Phase 1 usage before committing to native apps or advanced AI integrations.','Privacy by design: collect only information needed for the customer experience and protect project documents and personal data.'] : bullet(x)

doc.add_heading('Measures of success', level=1)
table(['Measure','What it tells us'],[
('Customer access success','Whether customers can reliably sign in and reach their project.'),
('Project information engagement','Whether customers view milestones, updates, and shared documents.'),
('Support request completion','Whether customers can ask for help and receive a response through the agreed workflow.'),
('Planning brief completion','In later phases, whether prospective customers can create and share a useful brief.'),
('AI usefulness and correction rate','In later phases, whether AI answers help customers and how often staff must correct them.')],[2.0,5.0])

doc.add_heading('Dependencies and decisions', level=1)
for x in ['Confirm the first release format: responsive web app/PWA is the recommended starting point.','Confirm how customer identity maps to project records and whether a customer may have multiple projects.','Agree which documents and progress updates are customer-visible, and who publishes them.','Choose the Phase 1 notification channel and define support ownership/response expectations.','Set delivery dates and budget after technical discovery and team sizing; this proposal does not assume either.'] : bullet(x)
callout('RECOMMENDED NEXT STEP','Approve the Phase 1 scope and confirm the app format, customer-to-project access rules, and update/notification workflow. Then convert the Phase 1 slices into implementation tickets and a delivery estimate.')

# Keep headings with following text and tables from splitting where possible
for p in doc.paragraphs:
 if p.style.name.startswith('Heading'):
  p.paragraph_format.keep_with_next=True
for t in doc.tables:
 for row in t.rows:
  trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)

doc.core_properties.title='ConstructONS Customer App Proposal and Phase Plan'
doc.core_properties.subject='Product proposal with Phase 1 project management scope and phased roadmap'
doc.core_properties.author='ConstructONS'
doc.save(out)
print(out)
